#!/usr/bin/env python3
"""Download Sumsub report PDFs, upload them to Google Drive, and update Excel.

The script intentionally does not collect or store Sumsub or Google passwords.
On first use, the browser opens for the user to sign in to Sumsub and Google.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import tempfile
from collections import OrderedDict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.http import MediaFileUpload
from openpyxl import load_workbook
from playwright.sync_api import Page, TimeoutError as PlaywrightTimeoutError, sync_playwright


SCOPES = ["https://www.googleapis.com/auth/drive.file"]
REQUIRED_HEADERS = {
    "id": "sumsub id",
    "sumsub_url": "sumsub_url",
    "drive_url": "google drive pdf url",
}
STATUS_HEADER = "PDF Processing Status (interim)"
GET_PDF_RE = re.compile(r"get\s*pdf", re.IGNORECASE)


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def normalise_header(value: Any) -> str:
    """Allow headers that differ only by spaces, underscores, or letter case."""
    return re.sub(r"[\s_]+", "", str(value or "").strip().casefold())


def safe_filename(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip())
    return cleaned[:120] or "sumsub_report"


def atomic_write_json(path: Path, content: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        "w", encoding="utf-8", dir=path.parent, delete=False, suffix=".tmp"
    ) as handle:
        json.dump(content, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
        temporary_path = Path(handle.name)
    os.replace(temporary_path, path)


def load_manifest(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {"schema_version": 1, "items": {}}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"Manifest is not valid JSON: {path}") from exc
    if not isinstance(data, dict) or not isinstance(data.get("items"), dict):
        raise RuntimeError(f"Manifest has an unexpected structure: {path}")
    return data


def get_drive_service(credentials_path: Path, token_path: Path):
    credentials: Credentials | None = None
    if token_path.exists():
        credentials = Credentials.from_authorized_user_file(str(token_path), SCOPES)
    if credentials and credentials.expired and credentials.refresh_token:
        credentials.refresh(Request())
    if not credentials or not credentials.valid:
        if not credentials_path.exists():
            raise RuntimeError(
                "Google OAuth credentials file was not found. See README_中文.md, step 2."
            )
        flow = InstalledAppFlow.from_client_secrets_file(str(credentials_path), SCOPES)
        credentials = flow.run_local_server(port=0)
    token_path.write_text(credentials.to_json(), encoding="utf-8")
    return build("drive", "v3", credentials=credentials, cache_discovery=False)


def drive_url(file_id: str) -> str:
    return f"https://drive.google.com/file/d/{file_id}/view?usp=drive_link"


def upload_pdf(
    service: Any,
    pdf_path: Path,
    folder_id: str,
    share_anyone_with_link: bool,
) -> tuple[str, str]:
    metadata = {"name": pdf_path.name, "parents": [folder_id]}
    media = MediaFileUpload(str(pdf_path), mimetype="application/pdf", resumable=True)
    created = (
        service.files()
        .create(
            body=metadata,
            media_body=media,
            fields="id,webViewLink",
            supportsAllDrives=True,
        )
        .execute()
    )
    file_id = created["id"]
    if share_anyone_with_link:
        (
            service.permissions()
            .create(
                fileId=file_id,
                body={"type": "anyone", "role": "reader", "allowFileDiscovery": False},
                fields="id",
                supportsAllDrives=True,
            )
            .execute()
        )
    return file_id, created.get("webViewLink") or drive_url(file_id)


def find_headers(sheet: Any) -> dict[str, int]:
    found: dict[str, int] = {}
    aliases = {key: normalise_header(label) for key, label in REQUIRED_HEADERS.items()}
    for column in range(1, sheet.max_column + 1):
        header = normalise_header(sheet.cell(1, column).value)
        for key, target in aliases.items():
            if header == target:
                found[key] = column
    missing = [REQUIRED_HEADERS[key] for key in REQUIRED_HEADERS if key not in found]
    if missing:
        raise RuntimeError("Missing required Excel column(s): " + ", ".join(missing))
    return found


def find_or_create_status_column(sheet: Any) -> int:
    target = normalise_header(STATUS_HEADER)
    for column in range(1, sheet.max_column + 1):
        if normalise_header(sheet.cell(1, column).value) == target:
            return column
    column = sheet.max_column + 1
    sheet.cell(1, column).value = STATUS_HEADER
    return column


def build_row_index(sheet: Any, columns: dict[str, int]) -> OrderedDict[str, list[int]]:
    rows: OrderedDict[str, list[int]] = OrderedDict()
    for row in range(2, sheet.max_row + 1):
        applicant_id = str(sheet.cell(row, columns["id"]).value or "").strip()
        source_url = str(sheet.cell(row, columns["sumsub_url"]).value or "").strip()
        if not applicant_id and not source_url:
            continue
        if not applicant_id or not source_url:
            print(f"Row {row}: skipped because sumsub ID or Sumsub_Url is blank.")
            continue
        rows.setdefault(applicant_id, []).append(row)
    return rows


def existing_drive_url(sheet: Any, rows: list[int], column: int) -> str | None:
    links = []
    for row in rows:
        value = str(sheet.cell(row, column).value or "").strip()
        if value and value.casefold() != "not found":
            links.append(value)
    if not links:
        return None
    if len(set(links)) > 1:
        print(f"Warning: duplicate Sumsub ID has different existing Drive URLs; keeping {links[0]}")
    return links[0]


def set_row_values(sheet: Any, rows: list[int], drive_column: int, status_column: int, url: str, status: str) -> None:
    for row in rows:
        cell = sheet.cell(row, drive_column)
        cell.value = url
        cell.hyperlink = url
        sheet.cell(row, status_column).value = status


def set_row_status(sheet: Any, rows: list[int], status_column: int, status: str) -> None:
    for row in rows:
        sheet.cell(row, status_column).value = status


def download_sumsub_pdf(page: Page, source_url: str, destination: Path) -> None:
    page.goto(source_url, wait_until="domcontentloaded", timeout=90_000)
    try:
        button = page.get_by_role("button", name=GET_PDF_RE)
        button.wait_for(state="visible", timeout=120_000)
    except PlaywrightTimeoutError as exc:
        raise RuntimeError(
            "Cannot find the Get PDF button. Log in to Sumsub in the opened browser, "
            "then rerun this command. If the report was deleted or unavailable, mark it not found manually."
        ) from exc
    try:
        with page.expect_download(timeout=120_000) as download_info:
            button.click()
        download = download_info.value
        destination.parent.mkdir(parents=True, exist_ok=True)
        download.save_as(str(destination))
    except PlaywrightTimeoutError as exc:
        raise RuntimeError("The Get PDF action did not start a download within two minutes.") from exc
    if not destination.exists() or destination.stat().st_size == 0:
        raise RuntimeError("Sumsub download did not create a usable PDF file.")


def choose_source_url(sheet: Any, rows: list[int], column: int) -> str:
    return str(sheet.cell(rows[0], column).value or "").strip()


def resolve_pdf_path(record: dict[str, Any], applicant_id: str, downloads_dir: Path) -> Path:
    """Use our deterministic filename first, then recognise a manually downloaded Sumsub PDF."""
    stored = str(record.get("pdf_path") or "").strip()
    if stored:
        candidate = Path(stored).expanduser()
        if candidate.exists():
            return candidate.resolve()
    deterministic = (downloads_dir / f"{safe_filename(applicant_id)}.pdf").resolve()
    if deterministic.exists():
        return deterministic
    if downloads_dir.exists():
        matches = [
            file for file in downloads_dir.iterdir()
            if file.is_file()
            and file.suffix.casefold() == ".pdf"
            and applicant_id.casefold() in file.name.casefold()
        ]
        if matches:
            # The newest is normally the correct one if Sumsub has generated the same report twice.
            return max(matches, key=lambda file: file.stat().st_mtime).resolve()
    return deterministic


def save_workbook(workbook: Any, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Download Sumsub PDFs, upload them to Google Drive, and fill an Excel column."
    )
    parser.add_argument("--input", required=True, type=Path, help="Input .xlsx workbook")
    parser.add_argument("--output", required=True, type=Path, help="Output .xlsx workbook")
    parser.add_argument("--sheet", help="Worksheet name. If omitted, the active worksheet is used.")
    parser.add_argument("--folder-id", required=True, help="Destination Google Drive folder ID")
    parser.add_argument(
        "--credentials",
        type=Path,
        required=True,
        help="Google OAuth desktop-app credentials JSON downloaded from Google Cloud",
    )
    parser.add_argument(
        "--token", type=Path, default=Path("token.json"), help="Google OAuth token cache path"
    )
    parser.add_argument(
        "--manifest",
        type=Path,
        default=Path("sumsub_drive_manifest.json"),
        help="Resume-state JSON path",
    )
    parser.add_argument(
        "--downloads-dir", type=Path, default=Path("pdf_downloads"), help="PDF download folder"
    )
    parser.add_argument(
        "--profile-dir",
        type=Path,
        default=Path("sumsub_browser_profile"),
        help="Persistent browser profile for Sumsub login session",
    )
    parser.add_argument(
        "--skip-download",
        action="store_true",
        help="Use existing local PDFs in --downloads-dir instead of visiting Sumsub",
    )
    parser.add_argument(
        "--share-anyone-with-link",
        action="store_true",
        help="Explicitly make each new PDF public to anyone with the link (use only if authorised).",
    )
    parser.add_argument("--headless", action="store_true", help="Run without displaying Sumsub browser")
    parser.add_argument("--max-records", type=int, help="Process at most this many unique Sumsub IDs")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_path = args.input.expanduser().resolve()
    output_path = args.output.expanduser().resolve()
    credentials_path = args.credentials.expanduser().resolve()
    token_path = args.token.expanduser().resolve()
    manifest_path = args.manifest.expanduser().resolve()
    downloads_dir = args.downloads_dir.expanduser().resolve()
    profile_dir = args.profile_dir.expanduser().resolve()

    if input_path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise RuntimeError("Input must be an .xlsx or .xlsm workbook.")
    if not input_path.exists():
        raise RuntimeError(f"Input workbook not found: {input_path}")
    if args.headless and not args.skip_download and not profile_dir.exists():
        raise RuntimeError("First Sumsub run must be visible. Remove --headless and sign in first.")

    workbook = load_workbook(input_path, keep_vba=input_path.suffix.casefold() == ".xlsm")
    if args.sheet:
        if args.sheet not in workbook.sheetnames:
            raise RuntimeError(f"Worksheet not found: {args.sheet}. Available: {', '.join(workbook.sheetnames)}")
        sheet = workbook[args.sheet]
    else:
        sheet = workbook.active
        print(f"No --sheet supplied; using active worksheet: {sheet.title}")

    columns = find_headers(sheet)
    status_column = find_or_create_status_column(sheet)
    indexed_rows = build_row_index(sheet, columns)
    manifest = load_manifest(manifest_path)
    items: dict[str, Any] = manifest["items"]
    service = get_drive_service(credentials_path, token_path)

    selected = list(indexed_rows.items())
    if args.max_records is not None:
        selected = selected[: max(args.max_records, 0)]

    processed = uploaded = retained = failures = 0
    browser_context = None
    playwright = None
    page = None
    try:
        if not args.skip_download:
            print("Opening a Sumsub browser. Complete any login or MFA there; do not close it while this script runs.")
            playwright = sync_playwright().start()
            browser_context = playwright.chromium.launch_persistent_context(
                str(profile_dir), headless=args.headless, accept_downloads=True
            )
            page = browser_context.pages[0] if browser_context.pages else browser_context.new_page()

        for applicant_id, rows in selected:
            processed += 1
            record = items.setdefault(applicant_id, {})
            existing_url = existing_drive_url(sheet, rows, columns["drive_url"])
            try:
                if existing_url:
                    record.update({"url": existing_url, "status": "existing_url", "updated_at": utc_now()})
                    set_row_values(sheet, rows, columns["drive_url"], status_column, existing_url, "Existing Google Drive URL retained")
                    retained += 1
                    atomic_write_json(manifest_path, manifest)
                    continue

                stored_url = str(record.get("url") or "").strip()
                if stored_url:
                    set_row_values(sheet, rows, columns["drive_url"], status_column, stored_url, "Recovered from resume manifest")
                    retained += 1
                    continue

                pdf_path = resolve_pdf_path(record, applicant_id, downloads_dir)
                if not pdf_path.exists():
                    if args.skip_download:
                        raise RuntimeError(f"Local PDF not found for --skip-download: {pdf_path}")
                    assert page is not None
                    source_url = choose_source_url(sheet, rows, columns["sumsub_url"])
                    print(f"[{processed}/{len(selected)}] Downloading PDF for {applicant_id} ...")
                    download_sumsub_pdf(page, source_url, pdf_path)
                    record.update({"pdf_path": str(pdf_path), "status": "downloaded", "updated_at": utc_now()})
                    atomic_write_json(manifest_path, manifest)

                existing_file_id = str(record.get("drive_file_id") or "").strip()
                if existing_file_id:
                    url = drive_url(existing_file_id)
                    status = "Recovered uploaded file from resume manifest"
                    retained += 1
                else:
                    print(f"[{processed}/{len(selected)}] Uploading {pdf_path.name} to Google Drive ...")
                    file_id, url = upload_pdf(service, pdf_path, args.folder_id, args.share_anyone_with_link)
                    record.update(
                        {
                            "pdf_path": str(pdf_path),
                            "drive_file_id": file_id,
                            "url": url,
                            "status": "uploaded",
                            "updated_at": utc_now(),
                        }
                    )
                    status = "Uploaded to Google Drive"
                    uploaded += 1

                set_row_values(sheet, rows, columns["drive_url"], status_column, url, status)
                atomic_write_json(manifest_path, manifest)
            except Exception as exc:  # Keep other rows running and leave a precise status in Excel.
                failures += 1
                message = str(exc).replace("\n", " ")[:300]
                record.update({"status": "error", "error": message, "updated_at": utc_now()})
                set_row_status(sheet, rows, status_column, f"ERROR: {message}")
                atomic_write_json(manifest_path, manifest)
                print(f"ERROR for {applicant_id}: {message}", file=sys.stderr)
            finally:
                save_workbook(workbook, output_path)
    finally:
        if browser_context:
            browser_context.close()
        if playwright:
            playwright.stop()

    print(
        f"Finished. Unique IDs processed: {processed}; uploaded: {uploaded}; "
        f"existing/resumed: {retained}; failed: {failures}.\nOutput: {output_path}"
    )
    return 1 if failures else 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("Interrupted. The manifest and latest output workbook have been saved; rerun the same command to resume.")
        raise SystemExit(130)
    except Exception as exc:
        print(f"Fatal error: {exc}", file=sys.stderr)
        raise SystemExit(2)
